"""Generate PPT portofolio dari outline. Jalankan: pip install python-pptx && python scripts/make_ppt.py"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

OUT = Path(__file__).resolve().parent.parent / "porto-ai-ppt.pptx"

SLIDES: list[tuple[str, list[str]]] = [
    ("OCR + RAG tanpa GPU",
     ["Scan kertas → tanya jawab berdasar dokumen (dengan sitasi)",
      "RapidOCR CPU + LangChain + pgvector + DeepSeek v4-flash",
      "github.com/adityanuriskandar17/ocr-rag-deepseek-pgvector"]),
    ("Masalah",
     ["Dokumen masih scan/foto → teks tidak searchable",
      "LLM saja: halusinasi + tidak tahu dokumen privat",
      "GPU mahal → butuh solusi CPU-friendly"]),
    ("Solusi",
     ["Scan → PyMuPDF 300 DPI → RapidOCR ONNX CPU",
      "Chunk 800/120 → embedding MiniLM lokal → pgvector",
      "Retriever top-5 → DeepSeek v4-flash → jawaban + sitasi"]),
    ("Arsitektur & Stack",
     ["RapidOCR (CPU): deteksi + recognition tanpa GPU",
      "MiniLM lokal: embedding ID/EN, hemat API",
      "pgvector :5433: vector store + metadata",
      "LangChain LCEL: retriever | prompt | llm | parser",
      "FastAPI + Streamlit: API & demo UI"]),
    ("Alur RAG",
     ["Ingest: OCR → chunk → embed → simpan (sekali)",
      "Retrieve: 5 chunk paling mirip (tiap tanya)",
      "Generate: hanya dari konteks + sitasi [hal. X]"]),
    ("Demo",
     ["Screenshot 1: upload scan → '42 chunks'",
      "Screenshot 2: tanya → jawaban + sources",
      "(Ganti dengan screenshot/GIF asli)"]),
    ("Hasil / Angka",
     ["Ingest 10 hal: __ mnt → __ mnt (paralel + batch)",
      "OCR: __ dtk/hal, conf __% | chunk: __",
      "Q&A selalu bersitasi (eval PASS)",
      "Isi angka dari log timing kamu"]),
    ("Tradeoff jujur",
     ["Bukan VLM GPU 3B: overkill untuk teks polos",
      "Lemah: tulisan tangan, stempel, tabel kompleks",
      "Mitigasi: ocr_conf + bbox + sitasi halaman"]),
    ("Roadmap",
     ["Reranker top-20 → top-5",
      "Hybrid BM25 + vektor",
      "Eval faithfulness + hit-rate",
      "LangGraph agent + observability"]),
    ("Penutup",
     ["github.com/adityanuriskandar17/ocr-rag-deepseek-pgvector",
      "Ingestion nyata > demo mainan; grounding > prompt panjang",
      "Terima kasih — Q&A"]),
]


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)  # 16:9
    title_layout = prs.slide_layouts[5]  # Title Only -> isi manual

    for i, (title, bullets) in enumerate(SLIDES):
        slide = prs.slides.add_slide(prs.slide_layouts[1] if i else prs.slide_layouts[0])
        if i == 0:
            slide.shapes.title.text = title
            slide.placeholders[1].text = "\n".join(bullets)
        else:
            slide.shapes.title.text = title
            tf = slide.placeholders[1].text_frame
            tf.clear()
            for j, b in enumerate(bullets):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = b
                p.level = 0
                for run in p.runs:
                    run.font.size = Pt(20)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
